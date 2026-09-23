#!/usr/bin/env python3
"""Generate a 16:9 PPTX of the FanDuel Help Portal experience assessment."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x45)
BLUE = RGBColor(0x00, 0x60, 0xA3)
INK = RGBColor(0x1A, 0x1A, 0x1A)
SLATE = RGBColor(0x5C, 0x6B, 0x73)
PAPER = RGBColor(0xF4, 0xF6, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x84, 0x4A)
CAUTION = RGBColor(0x8C, 0x4B, 0x02)

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, bold, color, font)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill_shape(s, color)
    return s


def content_chrome(slide, number, total=20):
    rect(slide, Inches(0), Inches(0), W, Inches(0.08), BLUE)
    add_textbox(
        slide,
        Inches(0.7),
        Inches(7.15),
        Inches(10),
        Inches(0.28),
        "FanDuel Help Portal  ·  Experience Assessment",
        size=11,
        color=SLATE,
    )
    add_textbox(
        slide,
        Inches(11.4),
        Inches(7.15),
        Inches(1.2),
        Inches(0.28),
        f"{number}  /  {total}",
        size=11,
        color=SLATE,
        align=PP_ALIGN.RIGHT,
    )


def title_bar(slide, title):
    add_textbox(slide, Inches(0.7), Inches(0.28), Inches(12), Inches(0.55), title, size=26, bold=True, color=NAVY)


def bullets(slide, l, t, w, h, items, size=16, color=INK, spacing=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(spacing)
        run = p.add_run()
        set_run(run, item, size=size, color=color)
    return box


def add_table(slide, l, t, w, h, headers, rows, col_widths=None):
    table_shape = slide.shapes.add_table(1 + len(rows), len(headers), l, t, w, h)
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = cw
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for run in p.runs:
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = WHITE
                run.font.name = "Calibri"
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.text = val
            if r % 2 == 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = PAPER
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(11)
                    run.font.color.rgb = INK
                    run.font.name = "Calibri"
                    run.font.bold = c == 0
    return table_shape


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def navy_slide(prs, kicker, title, subtitle=None, number=None):
    slide = blank(prs)
    rect(slide, Inches(0), Inches(0), W, H, NAVY)
    add_textbox(slide, Inches(0.8), Inches(2.2), Inches(11.5), Inches(0.4), kicker, size=13, bold=True, color=BLUE)
    add_textbox(slide, Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.4), title, size=36, bold=True, color=WHITE)
    if subtitle:
        add_textbox(slide, Inches(0.8), Inches(4.2), Inches(11.5), Inches(1.2), subtitle, size=18, color=RGBColor(0xC5, 0xD0, 0xDC))
    if number:
        add_textbox(slide, Inches(11.4), Inches(7.05), Inches(1.2), Inches(0.28), f"{number}  /  20", size=11, color=RGBColor(0x8A, 0x9B, 0xAB), align=PP_ALIGN.RIGHT)
    return slide


def phone(slide, l, t, lines, header="FanDuel Support"):
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(3.1), Inches(5.4))
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = RGBColor(0xC8, 0xD0, 0xD8)
    frame.line.width = Pt(1.25)
    hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, Inches(3.1), Inches(0.42))
    fill_shape(hdr, PAPER)
    hdr.line.fill.background()
    add_textbox(slide, l + Inches(0.12), t + Inches(0.08), Inches(2.86), Inches(0.28), header, size=11, bold=True, color=NAVY)
    y = t + Inches(0.55)
    for line in lines:
        add_textbox(slide, l + Inches(0.16), y, Inches(2.78), Inches(0.42), line, size=11, color=INK)
        y += Inches(0.38)


def card(slide, l, t, w, h, heading, body_lines):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = PAPER
    shp.line.fill.background()
    add_textbox(slide, l + Inches(0.22), t + Inches(0.16), w - Inches(0.4), Inches(0.4), heading, size=16, bold=True, color=BLUE)
    bullets(slide, l + Inches(0.22), t + Inches(0.55), w - Inches(0.4), h - Inches(0.7), body_lines, size=13, spacing=6)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 Cover
    navy_slide(
        prs,
        "EXPERIENCE ASSESSMENT  ·  SEPTEMBER 2026",
        "FanDuel Help Portal",
        "Digital Care\nPrepared for FanDuel  ·  Advisory + Design",
        1,
    )

    # 2 What this is
    s = blank(prs)
    content_chrome(s, 2)
    title_bar(s, "What this is")
    card(
        s,
        Inches(0.7),
        Inches(1.1),
        Inches(5.7),
        Inches(4.6),
        "This assessment",
        [
            "Current-state diagnosis of the help experience",
            "Design direction: homepage, articles, IA, agent",
            "How we should choose Agentforce hosting",
        ],
    )
    card(
        s,
        Inches(6.8),
        Inches(1.1),
        Inches(5.7),
        Inches(4.6),
        "This is not",
        [
            "A cost or savings case",
            "A delivery roadmap",
            "A recommendation to migrate to Casey",
        ],
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(5.9),
        Inches(11.8),
        Inches(0.9),
        "The chatbot should become an Agent. Where that agent lives depends on use cases we are still validating.",
        size=16,
        bold=True,
        color=NAVY,
    )

    # 3 Audience
    s = blank(prs)
    content_chrome(s, 3)
    title_bar(s, "The help site is an extension of the app")
    stats = [
        ("5.76M", "Site visitors"),
        ("1.80M", "Chats"),
        ("31.3%", "Overall chat rate"),
        ("82%", "Visits from a FanDuel app"),
    ]
    for i, (val, label) in enumerate(stats):
        x = Inches(0.7) + i * Inches(3.1)
        box = slide_stat = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.15), Inches(2.9), Inches(2.15))
        fill_shape(slide_stat, PAPER)
        add_textbox(s, x, Inches(1.4), Inches(2.9), Inches(0.9), val, size=32, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
        add_textbox(s, x, Inches(2.35), Inches(2.9), Inches(0.6), label, size=14, color=SLATE, align=PP_ALIGN.CENTER)
    add_textbox(
        s,
        Inches(0.7),
        Inches(3.55),
        Inches(12),
        Inches(0.4),
        "91% mobile or tablet   ·   ~275k external / organic visits (2025)",
        size=14,
        color=SLATE,
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(4.2),
        Inches(12),
        Inches(2.2),
        "Four in five users are already in the product. The site treats every visitor the same — no app context, no personalisation, no continuity.",
        size=20,
        bold=True,
        color=NAVY,
    )

    # 4 Five findings
    s = blank(prs)
    content_chrome(s, 4)
    title_bar(s, "Five findings")
    findings = [
        ("1  Search is buried", "Five category tiles above the fold; search is a header icon."),
        ("2  No app context is passed", "Casino and DFS see the same homepage. Chat starts at zero."),
        ("3  Articles bleed to agents", "Two different jobs — self-service format vs Necessary Transfer routing."),
        ("4  Taxonomy is product-led", "16 categories follow FanDuel’s org chart, not user intent."),
        ("5  Chat handoff has no continuity", "Customers repeat themselves. Agents start cold."),
    ]
    for i, (h, b) in enumerate(findings):
        y = Inches(1.05) + i * Inches(1.05)
        add_textbox(s, Inches(0.7), y, Inches(12), Inches(0.35), h, size=18, bold=True, color=BLUE)
        add_textbox(s, Inches(0.7), y + Inches(0.35), Inches(12), Inches(0.5), b, size=16, color=INK)

    # 5 Two jobs table
    s = blank(prs)
    content_chrome(s, 5)
    title_bar(s, "Chat rate is two jobs")
    add_textbox(
        s,
        Inches(0.7),
        Inches(0.85),
        Inches(12),
        Inches(0.35),
        "Do not read this as one deflection problem.",
        size=14,
        color=SLATE,
    )
    add_table(
        s,
        Inches(0.7),
        Inches(1.25),
        Inches(11.9),
        Inches(4.55),
        ["Article", "Visitors", "Chat rate", "Class"],
        [
            ["Location Troubleshooting Tips", "673,000", "5.5%", "Performing"],
            ["Why Am I Unable to Verify", "108,000", "22.6%", "Self-service"],
            ["Password Reset", "44,000", "43.4%", "Self-service"],
            ["Change My Password", "—", "39.1%", "Self-service"],
            ["2FA", "—", "23.1%", "Self-service"],
            ["How to Contact FanDuel Support", "—", "48.5%", "Used as a contact path"],
            ["Account Locked", "—", "62.6%", "Necessary Transfer"],
            ["Contact Us", "—", "61.8%", "Necessary Transfer"],
        ],
        col_widths=[Inches(5.2), Inches(1.8), Inches(1.8), Inches(3.1)],
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(5.95),
        Inches(11.9),
        Inches(0.95),
        "Location works because of an app deeplink. Password / verify / 2FA need steps. Account Locked and Contact Us need routing — not a how-to rewrite.",
        size=13,
        color=SLATE,
    )

    # 6 Current homepage
    s = blank(prs)
    content_chrome(s, 6)
    title_bar(s, "Current homepage — category first, search last")
    phone(
        s,
        Inches(0.8),
        Inches(1.05),
        [
            "Search icon only — buried",
            "My Account",
            "Manage My Money",
            "How to Play",
            "Rewards & Promos",
            "State Rules",
            "All-Star FAQs below fold",
            "Chat = small floating button",
        ],
    )
    bullets(
        s,
        Inches(4.4),
        Inches(1.4),
        Inches(8),
        Inches(5),
        [
            "Search is an icon in the header, not the way in.",
            "Five tiles force product-category navigation before intent.",
            "No trending or personalised content.",
            "Chat is an afterthought, not a service path.",
            "A Casino player and a DFS player see the same page.",
        ],
        size=18,
        spacing=14,
    )

    # 7 Opportunity
    s = blank(prs)
    content_chrome(s, 7)
    title_bar(s, "The opportunity is continuity")
    add_textbox(
        s,
        Inches(0.7),
        Inches(1.1),
        Inches(12),
        Inches(1.5),
        "82% arrive from a FanDuel app with a product, a session, and often a failed action. Today they leave that context, land on a generic homepage, and repeat themselves to a chatbot that starts at zero.",
        size=18,
        color=INK,
    )
    gaps = [
        ("1", "Self-service does not match how people arrive."),
        ("2", "Necessary Transfer is mixed with avoidable contact."),
        ("3", "The agent surface is disconnected from the product."),
    ]
    for i, (n, t) in enumerate(gaps):
        x = Inches(0.7) + i * Inches(4.05)
        c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.85), Inches(3.85), Inches(2.2))
        fill_shape(c, PAPER)
        add_textbox(s, x + Inches(0.2), Inches(3.05), Inches(3.4), Inches(0.4), n, size=22, bold=True, color=BLUE)
        add_textbox(s, x + Inches(0.2), Inches(3.5), Inches(3.45), Inches(1.3), t, size=16, color=NAVY)
    add_textbox(
        s,
        Inches(0.7),
        Inches(5.3),
        Inches(12),
        Inches(1.4),
        "Same intents, same content, and — if architecture allows — the same agent from app and Help Portal, with context passed in.",
        size=16,
        bold=True,
        color=NAVY,
    )

    # 8 Do either way
    s = blank(prs)
    content_chrome(s, 8)
    title_bar(s, "Do this either way — it does not wait on Aura vs Casey")
    add_table(
        s,
        Inches(0.7),
        Inches(1.15),
        Inches(11.9),
        Inches(4.4),
        ["Work", "Why it is independent of platform"],
        [
            ["Search as the homepage hero", "Fixes buried search on any Experience Cloud surface"],
            ["Top-10 article rewrites", "Format: steps, trees, inline actions"],
            ["Location-style deeplinks", "5.5% chat article already proves the pattern"],
            ["Intent-led IA (16 → 5)", "Taxonomy and redirects; still authored either way"],
            ["App → help context contract", "Continuity is a parameter contract, not a Casey feature"],
        ],
        col_widths=[Inches(4.4), Inches(7.5)],
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(5.8),
        Inches(12),
        Inches(0.8),
        "An agent cannot resolve what the content cannot express in steps.",
        size=16,
        bold=True,
        color=NAVY,
    )

    # 9 Strategic call
    s = blank(prs)
    content_chrome(s, 9)
    title_bar(s, "The chatbot should become an Agent")
    add_textbox(
        s,
        Inches(0.7),
        Inches(0.9),
        Inches(12),
        Inches(0.4),
        "That is the strategic call. Hosting is still open. We are not recommending a platform in this assessment.",
        size=15,
        color=SLATE,
    )
    card(
        s,
        Inches(0.7),
        Inches(1.45),
        Inches(5.85),
        Inches(4.9),
        "A  —  Independent Agentforce on current EC",
        [
            "Stay on Experience Cloud (Aura)",
            "Agentforce on the help site that exists today",
            "Can still do: search hero, step articles, deeplinks, intent IA, URL context",
            "App surface still needs a contract with app teams",
            "CMS: Salesforce-as-CMS constraints remain",
        ],
    )
    card(
        s,
        Inches(6.75),
        Inches(1.45),
        Inches(5.85),
        Inches(4.9),
        "B  —  Casey Help Portal",
        [
            "LWR Help Portal: site + agent as one product",
            "Dedicated content layer (CMS friction)",
            "In-app is an advantage only if the portal agent is the same agent the apps would open",
            "Do not treat in-app embed as a Casey feature until that validates",
        ],
    )

    # 10 Deciding question
    s = blank(prs)
    content_chrome(s, 10)
    title_bar(s, "The question we are validating with Agent teams")
    q = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.2), Inches(11.9), Inches(2.15))
    fill_shape(q, PAPER)
    add_textbox(
        s,
        Inches(0.95),
        Inches(1.45),
        Inches(11.4),
        Inches(1.7),
        "If the FanDuel app opens an agent, is that the same agent as the Help Portal agent — and can that agent orchestrate subagents (SBK, Casino, DFS, RG)?",
        size=20,
        bold=True,
        color=NAVY,
    )
    answers = [
        ("If yes", "Casey is the stronger hosting fit — one agent definition, one content layer, multiple surfaces."),
        ("If no", "Independent Agentforce on the current site may be cleaner. In-app is a separate agent."),
        ("If unknown", "Do not bind the experience work to Casey. Keep A and B open."),
    ]
    for i, (h, b) in enumerate(answers):
        y = Inches(3.55) + i * Inches(1.05)
        add_textbox(s, Inches(0.7), y, Inches(2.2), Inches(0.85), h, size=16, bold=True, color=BLUE)
        add_textbox(s, Inches(3.0), y, Inches(9.5), Inches(0.95), b, size=16, color=INK)

    # 11 Recs
    s = blank(prs)
    content_chrome(s, 11)
    title_bar(s, "Experience recommendations  ·  independent of A vs B")
    recs = [
        "1   Make search the homepage. Category tiles become secondary to intent.",
        "2   Rewrite high-bleed self-service articles into steps, trees, and deep links. Necessary Transfer is routing, not a how-to.",
        "3   Lead IA with user intent, not product lines. Replicate the Location deeplink across Login, Deposit, Withdrawal, Verify.",
        "4   Upgrade the chatbot to an Agent. Choose independent vs Casey once the same-agent question is answered.",
        "5   Pass app context (product, user_id, account_status, last_page) into whatever help or agent surface opens.",
    ]
    bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.4), recs, size=18, spacing=16)

    # 12 Proposed homepage
    s = blank(prs)
    content_chrome(s, 12)
    title_bar(s, "Proposed homepage — search first (either path)")
    phone(
        s,
        Inches(0.8),
        Inches(1.05),
        [
            "How can we help?",
            "Search as the hero",
            "Quick actions (app-aware)",
            "Location   ·   Deposit",
            "Withdrawal   ·   Login",
            "Trending (Amplitude)",
            "Reset password",
            "Withdrawal timelines",
            "Verify my account",
            "Chat with us",
        ],
    )
    bullets(
        s,
        Inches(4.4),
        Inches(1.4),
        Inches(8),
        Inches(5),
        [
            "This is the experience recommendation on Aura or Casey — not a Casey-only design.",
            "Search is the way in.",
            "Quick actions can use app source (product).",
            "Trending is Amplitude-driven, not a static FAQ list.",
            "Chat is a clear CTA. No invented wait-time SLA.",
        ],
        size=18,
        spacing=14,
    )

    # 13 Article
    s = blank(prs)
    content_chrome(s, 13)
    title_bar(s, "Articles as steps — then chat")
    phone(
        s,
        Inches(0.8),
        Inches(1.05),
        [
            "How to Reset Your Password",
            "1. Open the app, tap Log In",
            "2. Tap Forgot Password",
            "3. Enter email, Send Reset Link",
            "Reset password now  →",
            "",
            "Still stuck?",
            "Chat with us",
        ],
        header="Article",
    )
    bullets(
        s,
        Inches(4.4),
        Inches(1.4),
        Inches(8),
        Inches(5),
        [
            "Numbered steps replace narrative paragraphs.",
            "Inline deep link to the actual app action.",
            "Chat appears after the self-service path — not floating over the content.",
            "Same pattern for verify, 2FA, deposits, withdrawals.",
        ],
        size=18,
        spacing=14,
    )

    # 14 In-app
    s = blank(prs)
    content_chrome(s, 14)
    title_bar(s, "In-app agent is an experience target — not a Casey feature")
    phone(
        s,
        Inches(0.7),
        Inches(1.1),
        [
            "Bet Slip   ·   My Bets",
            "Promos   ·   Account",
            "",
            "Need help?",
            "Ask me anything about",
            "your account.",
            "Get help  →",
        ],
        header="Sportsbook  ·  Jordan · SBK",
    )
    phone(
        s,
        Inches(4.05),
        Inches(1.1),
        [
            "Logged in as Jordan · SBK",
            "Hi Jordan. I can see you’re",
            "on Sportsbook. What’s",
            "going on?",
            "",
            "Issue with a bet",
            "My account",
            "Deposit / withdrawal",
        ],
        header="Support  ·  context loaded",
    )
    add_textbox(
        s,
        Inches(7.45),
        Inches(1.3),
        Inches(5.2),
        Inches(4.8),
        "User never leaves the app.\n\nRequires the context contract — and, if validated, the same agent definition as the Help Portal rather than a second bot.\n\nPending Agent-team confirmation on same-agent and subagents.",
        size=16,
        color=INK,
    )

    # 15 IA
    s = blank(prs)
    content_chrome(s, 15)
    title_bar(s, "Intent-led IA  ·  16 product categories → 5 intents + Get Help")
    groups = [
        ("Log In & Access", "Reset · Locked · 2FA · Verify · Device"),
        ("Money In & Out", "Deposit · Withdrawal · Failed · Promos · Tax"),
        ("Bets, Games & Results", "Placed · Live · Disputes · Rules · Features"),
        ("Responsible Gaming", "Limits · Break · Reality checks · State RG"),
        ("State Rules & Location", "Where can I play · Location check · Rules"),
        ("Get Help", "Agentforce · Talk to a person · Complaint"),
    ]
    for i, (h, b) in enumerate(groups):
        col = i % 3
        row = i // 3
        x = Inches(0.7) + col * Inches(4.1)
        y = Inches(1.2) + row * Inches(2.5)
        c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.9), Inches(2.25))
        fill_shape(c, PAPER)
        add_textbox(s, x + Inches(0.2), y + Inches(0.25), Inches(3.5), Inches(0.7), h, size=16, bold=True, color=BLUE)
        add_textbox(s, x + Inches(0.2), y + Inches(1.0), Inches(3.5), Inches(1.0), b, size=14, color=INK)
    add_textbox(
        s,
        Inches(0.7),
        Inches(6.35),
        Inches(12),
        Inches(0.5),
        "Users think “I can’t withdraw,” not “I have a Casino question.”",
        size=15,
        color=SLATE,
    )

    # 16 Tiers
    s = blank(prs)
    content_chrome(s, 16)
    title_bar(s, "What the agent handles")
    tiers = [
        ("Tier 1  ·  Full self-service", GREEN, "Password reset, 2FA, deposit status, location check, verify account status.\n\nAgent gives steps + deep links."),
        ("Tier 2  ·  Guided + hand-off", CAUTION, "Withdrawal delays, failed transactions, bet dispute (info), promo eligibility.\n\nInfo first; escalate if unresolved."),
        ("Tier 3  ·  Human required", SLATE, "Account locked/suspended, identity fraud, regulatory complaint, RG escalation.\n\nAcknowledge, set expectations, route immediately."),
    ]
    for i, (h, col, b) in enumerate(tiers):
        x = Inches(0.7) + i * Inches(4.1)
        c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.2), Inches(3.9), Inches(4.4))
        fill_shape(c, PAPER)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.2), Inches(0.12), Inches(4.4))
        fill_shape(bar, col)
        add_textbox(s, x + Inches(0.3), Inches(1.4), Inches(3.4), Inches(0.9), h, size=16, bold=True, color=NAVY)
        add_textbox(s, x + Inches(0.3), Inches(2.4), Inches(3.4), Inches(2.8), b, size=14, color=INK)
    add_textbox(
        s,
        Inches(0.7),
        Inches(5.8),
        Inches(12),
        Inches(0.9),
        "Tier 3 is a routing problem. Tier 1 is a content and action problem. Do not design them as the same article pattern.",
        size=15,
        bold=True,
        color=NAVY,
    )

    # 17 RG
    s = blank(prs)
    content_chrome(s, 17)
    title_bar(s, "Responsible Gaming gate  —  hard rule")
    add_textbox(
        s,
        Inches(0.7),
        Inches(1.15),
        Inches(12),
        Inches(0.7),
        "If a session or account has any active RG flag (cooling off, self-exclusion, limit alert), the agent must:",
        size=16,
        color=INK,
    )
    rules = [
        "1   Acknowledge the user warmly and without judgement.",
        "2   Surface only RG support resources and state-mandated contacts.",
        "3   Not offer or discuss bonuses, promotions, or incentives in the same session.",
    ]
    bullets(s, Inches(0.7), Inches(2.0), Inches(12), Inches(2.4), rules, size=20, spacing=14)
    add_textbox(
        s,
        Inches(0.7),
        Inches(5.0),
        Inches(12),
        Inches(1.5),
        "Non-configurable. If subagents are supported, RG is a hard-gated subagent (or pre-agent policy), not a topic the general agent improvises.",
        size=16,
        bold=True,
        color=NAVY,
    )

    # 18 Context
    s = blank(prs)
    content_chrome(s, 18)
    title_bar(s, "Continuity lives in this contract")
    add_table(
        s,
        Inches(0.7),
        Inches(1.2),
        Inches(11.9),
        Inches(3.8),
        ["Parameter", "Source", "Use"],
        [
            ["product", "App identifier", "Product-relevant replies; route to subagent if allowed"],
            ["user_id", "Authenticated session", "Greet by name; skip account lookup"],
            ["account_status", "API on launch", "Proactive alert if the account has an issue"],
            ["last_page", "App navigation", "“I can see you were looking at withdrawals…”"],
        ],
        col_widths=[Inches(2.6), Inches(3.3), Inches(6.0)],
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(5.3),
        Inches(12),
        Inches(1.2),
        "Applies whether the surface is the current site, Casey, or an agent opened from the app.",
        size=16,
        bold=True,
        color=NAVY,
    )

    # 19 Validating
    s = blank(prs)
    content_chrome(s, 19)
    title_bar(s, "What we are validating next")
    nexts = [
        ("1", "Same agent in the app and on the Help Portal — and can it support subagents?", "Agent teams  ·  decides A vs B"),
        ("2", "Current chatbot platform / contract", "Digital Care"),
        ("3", "Which app teams own support entry in SBK / Casino / DFS", "The context contract needs owners"),
        ("4", "Amplitude / search data", "Trending and search-gap feedback"),
        ("5", "Content owner for 16 → 5 taxonomy", "Recommended either way"),
    ]
    for i, (n, t, who) in enumerate(nexts):
        y = Inches(1.1) + i * Inches(1.05)
        add_textbox(s, Inches(0.7), y, Inches(0.5), Inches(0.8), n, size=22, bold=True, color=BLUE)
        add_textbox(s, Inches(1.4), y, Inches(11.2), Inches(0.45), t, size=17, bold=True, color=NAVY)
        add_textbox(s, Inches(1.4), y + Inches(0.42), Inches(11.2), Inches(0.4), who, size=14, color=SLATE)

    # 20 Close
    navy_slide(
        prs,
        "FANDUEL HELP PORTAL  ·  DIGITAL CARE ASSESSMENT",
        "Experience first.\nAgent next.\nHosting when the architecture is clear.",
        "Do the experience work either way.  Upgrade the chatbot to an Agent.\nChoose independent vs Casey when we know whether the app opens the same agent — with subagents.",
        20,
    )

    out = "/Users/nsierchio/Documents/GitHub/fanduel-vip-console/FanDuel-Help-Portal-Experience-Assessment.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
